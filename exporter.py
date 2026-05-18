"""
exporter.py — запись результатов анализа в Excel.
"""

from pathlib import Path

import openpyxl


def append_row(xlsx_path: Path, tender_name: str, data: dict, fields: list):
    """
    Дописать строку в Excel.
    Создаёт файл с заголовком если не существует.
    Синхронизирует заголовок с текущим списком полей (если добавлены или переименованы).
    fields — список имён полей из config['fields'].
    """
    columns = ['файл'] + fields

    if xlsx_path.exists():
        wb = openpyxl.load_workbook(str(xlsx_path))
        ws = wb.active
        for i, col_name in enumerate(columns, start=1):
            if ws.cell(row=1, column=i).value != col_name:
                ws.cell(row=1, column=i, value=col_name)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(columns)

    row = [tender_name] + [data.get(col, '') for col in fields]
    ws.append(row)
    wb.save(str(xlsx_path))
    wb.close()
