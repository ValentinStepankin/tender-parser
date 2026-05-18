"""
extractor.py — распаковка архивов и извлечение текста из тендерных документов.
"""

import errno
import re
import subprocess
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

IMAGE_EXTS = frozenset({'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.tif', '.webp'})
SKIP_EXTS  = IMAGE_EXTS | frozenset({'.exe', '.dll', '.so', '.db', '.sqlite', '.lnk', '.ico', '.sys'})
ARCHIVE_EXTS = frozenset({'.zip', '.rar'})


# ---------------------------------------------------------------------------
# Публичный API
# ---------------------------------------------------------------------------

def extract_tender(source, errors_log_path=None) -> str:
    """
    Извлечь и склеить весь текст из одного тендера.
    source — путь к архиву (.zip/.rar) или одиночному файлу.
    Возвращает строку с разделителями по именам файлов.
    """
    source = Path(source)

    if source.stat().st_size == 0:
        return ""

    ext = source.suffix.lower()

    docs = []
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)

        if ext == '.zip':
            _extract_zip(source, tmpdir, errors_log_path)
            docs = _collect_docs(tmpdir, errors_log_path)
        elif ext == '.rar':
            _extract_rar(source, tmpdir, errors_log_path)
            docs = _collect_docs(tmpdir, errors_log_path)
        else:
            text = _extract_file_text(source, errors_log_path)
            if text:
                return f"=== {source.name} ===\n{text}"
            return ""

    docs.sort(key=lambda x: (0 if 'извещение' in x[0].lower() else 1, x[0].lower()))
    parts = [f"=== {name} ===\n{text}" for name, text in docs if text.strip()]
    return "\n\n".join(parts)


def extract_archive(archive_path, dest, errors_log_path=None):
    """Распаковать архив в папку dest."""
    path = Path(archive_path)
    ext = path.suffix.lower()
    if ext == '.zip':
        _extract_zip(path, Path(dest), errors_log_path)
    elif ext == '.rar':
        _extract_rar(path, Path(dest), errors_log_path)


def peek_archive_contents(archive_path) -> list:
    """Вернуть список имён файлов в архиве (без распаковки)."""
    path = Path(archive_path)
    ext = path.suffix.lower()
    try:
        if ext == '.zip':
            with zipfile.ZipFile(path) as zf:
                return [n for n in zf.namelist() if not n.endswith('/')]
        elif ext == '.rar':
            import rarfile
            with rarfile.RarFile(str(path)) as rf:
                return [n for n in rf.namelist() if not n.endswith('/')]
    except Exception:
        pass
    return []


def is_archive(path) -> bool:
    return Path(path).suffix.lower() in ARCHIVE_EXTS


# ---------------------------------------------------------------------------
# Сборка документов из директории
# ---------------------------------------------------------------------------

def _collect_docs(directory: Path, errors_log_path) -> list:
    """Рекурсивно собрать [(name, text)] из всех документов в директории."""
    result = []

    for item in sorted(directory.rglob('*'), key=lambda p: p.name.lower()):
        if not item.is_file():
            continue

        ext = item.suffix.lower()

        if ext in SKIP_EXTS:
            continue

        if ext == '.zip':
            try:
                with tempfile.TemporaryDirectory() as sub_tmp:
                    _extract_zip(item, Path(sub_tmp), errors_log_path)
                    sub_docs = _collect_docs(Path(sub_tmp), errors_log_path)
                result.extend(sub_docs)
            except OSError as e:
                if e.errno == errno.ENOSPC:
                    _log_error(errors_log_path, str(item), "Недостаточно места на диске — архив пропущен")
                else:
                    _log_error(errors_log_path, str(item), str(e))
            except Exception as e:
                _log_error(errors_log_path, str(item), str(e))
        elif ext == '.rar':
            try:
                with tempfile.TemporaryDirectory() as sub_tmp:
                    _extract_rar(item, Path(sub_tmp), errors_log_path)
                    sub_docs = _collect_docs(Path(sub_tmp), errors_log_path)
                result.extend(sub_docs)
            except OSError as e:
                if e.errno == errno.ENOSPC:
                    _log_error(errors_log_path, str(item), "Недостаточно места на диске — архив пропущен")
                else:
                    _log_error(errors_log_path, str(item), str(e))
            except Exception as e:
                _log_error(errors_log_path, str(item), str(e))
        else:
            text = _extract_file_text(item, errors_log_path)
            if text is not None:
                result.append((item.name, text))

    return result


# ---------------------------------------------------------------------------
# Распаковка архивов
# ---------------------------------------------------------------------------

def _extract_zip(zip_path: Path, dest: Path, errors_log_path):
    # Сначала пробуем zipfile — быстрый, корректно правит CP866-имена
    try:
        with zipfile.ZipFile(zip_path) as zf:
            for member in zf.infolist():
                # Исправить кодировку имён файлов: CP437 → CP866 для русских имён
                if member.flag_bits & 0x800:
                    filename = member.filename  # UTF-8, zipfile уже декодировал корректно
                else:
                    try:
                        filename = member.filename.encode('cp437').decode('cp866')
                    except (UnicodeEncodeError, UnicodeDecodeError):
                        filename = member.filename

                target = dest / filename
                if member.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(member) as src, open(target, 'wb') as dst:
                    dst.write(src.read())
        return
    except zipfile.BadZipFile:
        pass  # битый zip — пробуем unar
    except Exception as e:
        _log_error(errors_log_path, str(zip_path), str(e))
        return

    # Fallback: unar — некоторые «битые» zip-архивы он всё-таки распаковывает
    try:
        result = subprocess.run(
            ['unar', '-o', str(dest), '-D', '-f', str(zip_path)],
            capture_output=True, timeout=120, check=False,
        )
        # Даже при ненулевом коде часть файлов могла извлечься — пусть _collect_docs их прочитает
        if result.returncode != 0:
            _log_error(errors_log_path, str(zip_path), f"ZIP частично распакован через unar (код {result.returncode})")
    except FileNotFoundError:
        _log_error(errors_log_path, str(zip_path), "Повреждённый ZIP, unar не установлен (brew install unar)")
    except Exception as e:
        _log_error(errors_log_path, str(zip_path), f"Повреждённый ZIP, unar упал: {e}")


def _extract_rar(rar_path: Path, dest: Path, errors_log_path):
    # unar (The Unarchiver) — подписан, работает на Apple Silicon без Gatekeeper-проблем
    try:
        result = subprocess.run(
            ['unar', '-o', str(dest), '-D', '-f', str(rar_path)],
            capture_output=True, timeout=120, check=False,
        )
        if result.returncode == 0:
            return
    except FileNotFoundError:
        pass  # unar не установлен — пробуем rarfile+unrar

    # Fallback: rarfile + unrar
    try:
        import rarfile
        with rarfile.RarFile(str(rar_path)) as rf:
            rf.extractall(str(dest))
    except ImportError:
        _log_error(errors_log_path, str(rar_path), "rarfile не установлен: pip install rarfile")
    except Exception as e:
        err = str(e)
        if 'password' in err.lower() or 'wrong password' in err.lower():
            _log_error(errors_log_path, str(rar_path), "Архив защищён паролем")
        elif 'cannot exec' in err.lower() or 'unrar' in err.lower():
            import sys
            hint = 'brew install unar' if sys.platform != 'win32' else 'установите unrar.exe с rarlab.com и добавьте в PATH'
            _log_error(errors_log_path, str(rar_path), f"unrar не установлен ({hint})")
        else:
            _log_error(errors_log_path, str(rar_path), f"Ошибка RAR: {e}")


# ---------------------------------------------------------------------------
# Извлечение текста из файлов
# ---------------------------------------------------------------------------

def _extract_file_text(path: Path, errors_log_path) -> str | None:
    """Извлечь текст. None = файл пропущен намеренно."""
    ext = path.suffix.lower()
    try:
        if ext == '.pdf':
            return _extract_pdf(path, errors_log_path)
        elif ext == '.docx':
            return _extract_docx(path)
        elif ext == '.doc':
            return _extract_doc(path, errors_log_path)
        elif ext == '.rtf':
            return _extract_rtf(path)
        elif ext in ('.xlsx', '.xls'):
            return _extract_spreadsheet(path)
        elif ext in ('.odt', '.ods'):
            return _extract_odf(path)
        elif ext in ('.html', '.htm'):
            return _extract_html_file(path)
        elif ext == '.xml':
            return _extract_xml(path)
        elif ext in ('.txt', '.csv', '.md'):
            return _read_text_file(path)
        else:
            return None
    except Exception as e:
        _log_error(errors_log_path, str(path), str(e))
        return ""


# --- PDF ---

def _extract_pdf(path: Path, errors_log_path) -> str:
    try:
        import fitz
    except ImportError:
        _log_error(errors_log_path, str(path), "PyMuPDF не установлен: pip install PyMuPDF")
        return ""

    doc = fitz.open(str(path))
    parts = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            parts.append(text)
    doc.close()

    result = "\n\n".join(parts)
    if len(result) < 50:
        _log_error(errors_log_path, str(path), "PDF-скан, текст не извлечён")
        return "[только сканы, текст не извлечён]"
    return result


# --- DOCX ---

def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


# --- DOC ---

def _extract_doc(path: Path, errors_log_path) -> str:
    """
    .doc в тендерных системах часто HTML-файл с расширением .doc.
    Определяем по заголовку: если начинается с <html — читаем как HTML.
    Иначе конвертируем через LibreOffice.
    """
    raw = _read_bytes(path)
    if not raw:
        return ""

    header = raw[:200].lstrip()
    if header.lower().startswith(b'<html') or header.lower().startswith(b'<!doctype'):
        return _extract_html_bytes(raw)

    return _convert_with_libreoffice(path, errors_log_path)


def _find_libreoffice() -> str | None:
    """Найти LibreOffice на текущей платформе."""
    import shutil, sys
    for cmd in ('libreoffice', 'soffice', 'soffice.exe'):
        if shutil.which(cmd):
            return cmd
    if sys.platform == 'win32':
        candidate = Path(r'C:\Program Files\LibreOffice\program\soffice.exe')
        if candidate.exists():
            return str(candidate)
    return None


def _convert_with_libreoffice(path: Path, errors_log_path) -> str:
    cmd = _find_libreoffice()
    if cmd is None:
        _log_error(errors_log_path, str(path), "LibreOffice не установлен, .doc пропущен")
        return ""

    with tempfile.TemporaryDirectory() as tmpdir:
        try:
            subprocess.run(
                [cmd, '--headless', '--convert-to', 'docx',
                 '--outdir', tmpdir, str(path)],
                capture_output=True, timeout=60, check=False,
            )
        except subprocess.TimeoutExpired:
            _log_error(errors_log_path, str(path), "LibreOffice timeout")
            return ""

        docx_files = list(Path(tmpdir).glob('*.docx'))
        if not docx_files:
            _log_error(errors_log_path, str(path), "LibreOffice не создал .docx")
            return ""

        return _extract_docx(docx_files[0])


# --- HTML ---

def _extract_html_file(path: Path) -> str:
    return _extract_html_bytes(_read_bytes(path))


def _extract_html_bytes(raw: bytes) -> str:
    from html.parser import HTMLParser

    encoding = _detect_encoding(raw)
    text = raw.decode(encoding, errors='replace')

    class _TextExtractor(HTMLParser):
        def __init__(self):
            super().__init__()
            self.parts = []
            self._skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ('script', 'style'):
                self._skip += 1

        def handle_endtag(self, tag):
            if tag in ('script', 'style'):
                self._skip = max(0, self._skip - 1)

        def handle_data(self, data):
            if not self._skip:
                s = data.strip()
                if s:
                    self.parts.append(s)

    parser = _TextExtractor()
    parser.feed(text)
    return ' '.join(parser.parts)


# --- RTF ---

def _extract_rtf(path: Path) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError:
        raise RuntimeError("striprtf не установлен: pip install striprtf")
    raw = _read_bytes(path)
    encoding = _detect_encoding(raw)
    return rtf_to_text(raw.decode(encoding, errors='replace'))


# --- Spreadsheets ---

def _extract_spreadsheet(path: Path) -> str:
    if path.suffix.lower() == '.xls':
        return _extract_xls(path)
    return _extract_xlsx(path)


def _extract_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row
                     if c is not None and str(c).strip() not in ('', 'None')]
            if cells:
                rows.append('\t'.join(cells))
        if rows:
            parts.append(f"[{sheet.title}]\n" + '\n'.join(rows))
    wb.close()
    return '\n\n'.join(parts)


def _extract_xls(path: Path) -> str:
    import xlrd
    wb = xlrd.open_workbook(str(path))
    parts = []
    for i in range(wb.nsheets):
        sheet = wb.sheet_by_index(i)
        rows = []
        for rx in range(sheet.nrows):
            cells = [str(sheet.cell(rx, cx).value).strip()
                     for cx in range(sheet.ncols)
                     if str(sheet.cell(rx, cx).value).strip() not in ('', 'None')]
            if cells:
                rows.append('\t'.join(cells))
        if rows:
            parts.append(f"[{sheet.name}]\n" + '\n'.join(rows))
    return '\n\n'.join(parts)


# --- ODF ---

def _extract_odf(path: Path) -> str:
    try:
        from odf import text as odf_text, teletype
        from odf.opendocument import load
    except ImportError:
        raise RuntimeError("odfpy не установлен: pip install odfpy")
    doc = load(str(path))
    return teletype.extractText(doc.text)


# --- XML ---

def _extract_xml(path: Path) -> str:
    import xml.etree.ElementTree as ET
    raw = _read_bytes(path)
    encoding = _detect_encoding(raw)
    text = raw.decode(encoding, errors='replace')
    try:
        root = ET.fromstring(text)
        return ' '.join(root.itertext())
    except ET.ParseError:
        return re.sub(r'<[^>]+>', ' ', text)


# --- Plain text ---

def _read_text_file(path: Path) -> str:
    raw = _read_bytes(path)
    encoding = _detect_encoding(raw)
    return raw.decode(encoding, errors='replace')


# ---------------------------------------------------------------------------
# Вспомогательные
# ---------------------------------------------------------------------------

def _read_bytes(path: Path) -> bytes:
    with open(path, 'rb') as f:
        return f.read()


def _detect_encoding(raw: bytes) -> str:
    if not raw:
        return 'utf-8'
    try:
        import chardet
        result = chardet.detect(raw[:4096])
        return result.get('encoding') or 'utf-8'
    except ImportError:
        return 'utf-8'


def _log_error(log_path, filename: str, reason: str):
    if not log_path:
        return
    log_path = Path(log_path)
    log_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with open(log_path, 'a', encoding='utf-8') as f:
        f.write(f"{timestamp} | {filename} | {reason}\n")
