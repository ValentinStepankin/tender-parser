"""
tender.py — точка входа, CLI, оркестрация.

Режимы:
  (без --mode)   — только извлечение текста → .txt
  --mode api     — анализ через OpenRouter / ChatGPT → Excel
  --mode ollama  — анализ через локальную модель → Excel
"""

import argparse
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import yaml
from dotenv import load_dotenv

from analyzer import APICallError, analyze
from exporter import append_row
from extractor import extract_archive, extract_tender, is_archive, peek_archive_contents


def main():
    args = _parse_args()
    load_dotenv()
    config = _load_config()

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        print(f"Ошибка: путь не существует: {input_path}", file=sys.stderr)
        sys.exit(1)

    errors_log = Path(__file__).parent / 'logs' / 'errors.log'
    errors_log.parent.mkdir(exist_ok=True)

    output_dir = _make_output_dir(input_path)

    with tempfile.TemporaryDirectory(prefix='tender_container_') as tmp:
        tmp_path = Path(tmp)
        tenders = _discover_tenders(input_path, tmp_path, errors_log)

        total = len(tenders)
        print(f"Найдено тендеров: {total}")
        print(f"Результаты:       {output_dir}")
        if args.mode:
            print(f"Режим анализа:    {args.mode}")
        print()

        xlsx_path = output_dir / 'result.xlsx'
        processed = _load_processed(xlsx_path)

        for i, (name, source) in enumerate(tenders, 1):
            txt_path = output_dir / f"{name}.txt"

            # Читаем готовый .txt или извлекаем заново
            if txt_path.exists():
                text = txt_path.read_text(encoding='utf-8')
            else:
                if source is None:
                    print(f"[{i}/{total}] {name} — пропущен (нет источника)")
                    continue
                source_path = Path(source)
                if not source_path.exists() or source_path.stat().st_size == 0:
                    print(f"[{i}/{total}] {name} — пропущен (пустой архив)")
                    continue
                text = extract_tender(source, errors_log)
                if not text.strip():
                    print(f"[{i}/{total}] {name} — пропущен (нет текста)")
                    continue
                txt_path.write_text(text, encoding='utf-8')

            if args.mode is None:
                print(f"[{i}/{total}] {name} — текст извлечён")
                continue

            if name in processed:
                print(f"[{i}/{total}] {name} — пропущен (уже в Excel)")
                continue

            # Анализ через LLM
            try:
                data = analyze(text, config, args.mode)
            except APICallError as e:
                _log_error(errors_log, name, f"Ошибка API: {e}")
                print(f"\n[{i}/{total}] {name} — ошибка API, прогон остановлен:")
                print(f"  {e}\n")
                break
            except Exception as e:
                _log_error(errors_log, name, f"Ошибка: {e}")
                print(f"[{i}/{total}] {name} — пропущен (ошибка: {e})")
                continue

            filled = sum(1 for v in data.values() if v and str(v).strip())
            if filled < 3:
                _log_error(errors_log, name, f"Модель вернула менее 3 полей ({filled})")
                print(f"[{i}/{total}] {name} — пропущен (модель вернула {filled} поля(ей))")
                continue

            print(f"[{i}/{total}] {name} — готово")
            fields = [f['name'] for f in config.get('fields', [])]
            append_row(xlsx_path, name, data, fields)

    if args.mode:
        print(f"\nГотово. Excel: {xlsx_path}")


# ---------------------------------------------------------------------------
# Определение тендеров
# ---------------------------------------------------------------------------

def _discover_tenders(input_path: Path, tmp_dir: Path, errors_log: Path) -> list:
    """
    Вернуть список [(tender_name, source_path_or_None)].
    source_path=None означает что .txt уже должен существовать.
    """
    if input_path.is_dir():
        return _tenders_from_dir(input_path)

    if is_archive(input_path):
        contents = peek_archive_contents(input_path)
        inner_archives = [f for f in contents
                          if Path(f).suffix.lower() in {'.zip', '.rar'}]

        if inner_archives:
            # Контейнер: распаковываем один раз в tmp, каждый вложенный архив = тендер
            extract_archive(input_path, tmp_dir, errors_log)
            tenders = []
            for item in sorted(tmp_dir.rglob('*'), key=lambda p: p.name.lower()):
                if item.is_file() and item.suffix.lower() in {'.zip', '.rar'}:
                    tenders.append((item.stem, str(item)))
            return tenders

        # Сам архив = один тендер
        return [(input_path.stem, str(input_path))]

    # Одиночный файл = один тендер
    return [(input_path.stem, str(input_path))]


def _tenders_from_dir(directory: Path) -> list:
    """Найти тендеры в папке: архивы и одиночные документы."""
    doc_exts = {'.doc', '.docx', '.pdf', '.rtf', '.odt'}
    tenders = []
    for item in sorted(directory.iterdir(), key=lambda p: p.name.lower()):
        if not item.is_file():
            continue
        ext = item.suffix.lower()
        if ext in {'.zip', '.rar'} or ext in doc_exts:
            tenders.append((item.stem, str(item)))
    return tenders


# ---------------------------------------------------------------------------
# Вспомогательные
# ---------------------------------------------------------------------------

def _make_output_dir(input_path: Path) -> Path:
    base = input_path.parent / f"{input_path.stem}_parsed"
    base.mkdir(exist_ok=True)
    return base


def _load_config() -> dict:
    config_path = Path(__file__).parent / 'config.yaml'
    with open(config_path, 'r', encoding='utf-8') as f:
        return yaml.safe_load(f)


def _parse_args():
    parser = argparse.ArgumentParser(description='Анализатор тендерной документации')
    parser.add_argument('--input', required=True,
                        help='Путь к архиву или папке с тендерами')
    parser.add_argument('--mode', choices=['api', 'ollama'], default=None,
                        help='Режим анализа: api или ollama')
    return parser.parse_args()


def _load_processed(xlsx_path: Path) -> set:
    """Вернуть множество имён тендеров уже записанных в Excel."""
    if not xlsx_path.exists():
        return set()
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True)
    ws = wb.active
    names = {ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=1).value}
    wb.close()
    return names


def _log_error(log_path: Path, filename: str, reason: str):
    log_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with open(log_path, 'a', encoding='utf-8') as f:
        f.write(f"{timestamp} | {filename} | {reason}\n")


if __name__ == '__main__':
    main()
